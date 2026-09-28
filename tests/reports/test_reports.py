import io
from datetime import date, timedelta

from httpx2 import AsyncClient
from openpyxl import load_workbook

from src.reports import exporters
from src.users.models import UserRole
from tests.helpers import login, make_user
from tests.tickets.helpers import make_ticket

DATE_RANGE = {"start": (date.today() - timedelta(days=30)).isoformat(), "end": date.today().isoformat()}


async def test_reports_require_agent(client: AsyncClient, db) -> None:
    await make_user(db, "u@example.com")
    await login(client, "u@example.com")
    assert (await client.get("/reports")).status_code == 403


async def test_reports_page_and_partials(client: AsyncClient, db) -> None:
    await make_user(db, "agent@example.com", UserRole.AGENT)
    requester = await make_user(db, "u@example.com")
    await make_ticket(db, requester, subject="Report ticket")
    await login(client, "agent@example.com")

    assert (await client.get("/reports")).status_code == 200
    for key in ("ticket-volume", "status", "priority", "sla", "agents", "kb"):
        resp = await client.get(f"/reports/partials/{key}", params=DATE_RANGE)
        assert resp.status_code == 200
    assert (await client.get("/reports/partials/unknown")).status_code == 400


async def test_report_exports(client: AsyncClient, db) -> None:
    await make_user(db, "agent@example.com", UserRole.AGENT)
    requester = await make_user(db, "u@example.com")
    await make_ticket(db, requester)
    await login(client, "agent@example.com")

    expectations = {
        "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "pdf": "application/pdf",
    }
    for fmt, content_type in expectations.items():
        resp = await client.get("/reports/export", params={"report": "status", "format": fmt, **DATE_RANGE})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(content_type)
        assert "attachment" in resp.headers["content-disposition"]
        assert len(resp.content) > 0

    assert (await client.get("/reports/export", params={"report": "status", "format": "xml"})).status_code == 400


async def test_export_range_guard(client: AsyncClient, db) -> None:
    await make_user(db, "agent@example.com", UserRole.AGENT)
    await login(client, "agent@example.com")
    resp = await client.get("/reports/export", params={"report": "status", "format": "csv", "start": "2000-01-01", "end": "2020-01-01"})
    assert resp.status_code == 400


def test_exporter_builders() -> None:
    columns = ["Name", "Count"]
    rows = [["Open", "3"], ["Closed", "1"]]
    assert b"Open" in exporters.build_csv(columns, rows)
    assert exporters.build_xlsx("Report", columns, rows)[:2] == b"PK"
    assert exporters.build_pdf("Report", "subtitle", columns, rows)[:4] == b"%PDF"


def test_csv_neutralizes_formula_injection() -> None:
    payloads = ["=cmd|'/c calc'!A1", "+1+1", "-2+3", "@SUM(A1)", "\tstart", "\rcarriage"]
    text = exporters.build_csv(["Name"], [[payload] for payload in payloads]).decode("utf-8-sig")
    for payload in payloads:
        assert f"'{payload}" in text


def test_xlsx_stores_formula_as_text() -> None:
    data = exporters.build_xlsx("Report", ["Name"], [["=cmd|'/c calc'!A1"]])
    sheet = load_workbook(io.BytesIO(data), data_only=False).active
    cell = sheet["A2"]
    assert cell.data_type == "s"
    assert cell.value == "=cmd|'/c calc'!A1"


async def test_export_sanitizes_formula_agent_name(client: AsyncClient, db) -> None:
    agent = await make_user(db, "agent@example.com", UserRole.AGENT)
    agent.name = "=cmd|'/c calc'!A1"
    requester = await make_user(db, "u@example.com")
    ticket = await make_ticket(db, requester)
    ticket.assignee_id = agent.id
    await db.commit()
    await login(client, "agent@example.com")

    csv = await client.get("/reports/export", params={"report": "agents", "format": "csv", **DATE_RANGE})
    assert csv.status_code == 200
    assert "'=cmd|'/c calc'!A1" in csv.text

    xlsx = await client.get("/reports/export", params={"report": "agents", "format": "xlsx", **DATE_RANGE})
    assert xlsx.status_code == 200
    sheet = load_workbook(io.BytesIO(xlsx.content), data_only=False).active
    formula_cells = [cell for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=cmd")]
    assert formula_cells
    assert all(cell.data_type == "s" for cell in formula_cells)
